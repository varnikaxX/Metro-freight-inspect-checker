"""PDF and editable compliance-report generation for Sanchayak.

Install the optional report dependencies with:
    pip install reportlab python-docx
"""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import httpx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    Image as ReportLabImage,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

# Sanchayak design-system colours.
PARCHMENT = colors.HexColor("#F3EDE3")
INK = colors.HexColor("#3B2F26")
BURNT_ORANGE = colors.HexColor("#D9762B")
SAGE_GREEN = colors.HexColor("#5C7A5C")
BRICK_RED = colors.HexColor("#B4483A")
AMBER = colors.HexColor("#C9922A")
TABLE_GRID = colors.HexColor("#CBBBAA")
TABLE_ROW = colors.HexColor("#FBF7F0")

FIELD_ORDER = (
    "manufacturer_address",
    "generic_name",
    "net_quantity",
    "manufacture_date",
    "MRP",
    "consumer_care",
    "country_of_origin",
    "best_before_date",
)
FIELD_LABELS = {
    "manufacturer_address": "Manufacturer / packer / importer address",
    "generic_name": "Generic/common commodity name",
    "net_quantity": "Net quantity",
    "manufacture_date": "Manufacture / packing / import date",
    "MRP": "MRP (inclusive of taxes)",
    "consumer_care": "Consumer care / complaint details",
    "country_of_origin": "Country of origin",
    "best_before_date": "Best-before / expiry date",
}
VERDICT_COLOURS = {
    "COMPLIANT": SAGE_GREEN,
    "NON COMPLIANT": BRICK_RED,
    "NEEDS REVIEW": AMBER,
}
STATUS_COLOURS = {"pass": SAGE_GREEN, "fail": BRICK_RED, "flagged": AMBER}


def _display(value: Any, fallback: str) -> str:
    """Return a human-friendly representation for nullable report values."""
    return fallback if value is None or str(value).strip() == "" else str(value)


def _normalised_fields(scan_result: dict[str, Any]) -> list[dict[str, Any]]:
    """Return all mandatory fields in a fixed order, filling absent records safely."""
    by_name = {item.get("field_name"): item for item in scan_result.get("fields", [])}
    return [
        by_name.get(
            name,
            {
                "field_name": name,
                "extracted_value": None,
                "status": "flagged",
                "reason": "No result supplied by the compliance engine.",
            },
        )
        for name in FIELD_ORDER
    ]


def _status_colour(status: str) -> colors.Color:
    """Return a ReportLab ``Color`` for a status, including unknown statuses."""
    return STATUS_COLOURS.get(status.strip().lower(), AMBER)


def _read_image(source: str | None) -> BytesIO | None:
    """Load an image from a local path or HTTP(S) URL, returning ``None`` if unavailable."""
    if not source:
        return None
    try:
        parsed = urlparse(source)
        if parsed.scheme in {"http", "https"}:
            response = httpx.get(source, timeout=10.0, follow_redirects=True)
            response.raise_for_status()
            return BytesIO(response.content)
        path = Path(parsed.path if parsed.scheme == "file" else source)
        if path.is_file():
            return BytesIO(path.read_bytes())
    except (OSError, httpx.HTTPError, ValueError):
        pass
    return None


def _report_styles() -> dict[str, ParagraphStyle]:
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("SanchayakTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=19, leading=23, textColor=INK, alignment=TA_CENTER, spaceAfter=7),
        "meta": ParagraphStyle("Metadata", parent=styles["BodyText"], fontSize=9, leading=13, textColor=INK, alignment=TA_CENTER),
        "heading": ParagraphStyle("SectionHeading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=13, leading=17, textColor=BURNT_ORANGE, spaceBefore=12, spaceAfter=6),
        "body": ParagraphStyle("ReportBody", parent=styles["BodyText"], fontSize=8.5, leading=11, textColor=INK),
        "table_header": ParagraphStyle("TableHeader", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=PARCHMENT, alignment=TA_LEFT),
        "verdict": ParagraphStyle("Verdict", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=17, leading=22, textColor=colors.white, alignment=TA_CENTER),
        "caption": ParagraphStyle("Caption", parent=styles["BodyText"], fontName="Helvetica-Oblique", fontSize=8, leading=10, textColor=INK, alignment=TA_CENTER),
    }


def _image_story(source: str | None, caption: str, styles: dict[str, ParagraphStyle]) -> list[Any]:
    """Create a fitted image and caption, or a readable unavailable-image notice."""
    data = _read_image(source)
    if data is None:
        return [Paragraph(f"{caption}: <i>Image unavailable</i>", styles["body"]), Spacer(1, 0.2 * cm)]
    try:
        image = ReportLabImage(data)
        image._restrictSize(16.5 * cm, 10.5 * cm)
        return [KeepTogether([image, Spacer(1, 0.08 * cm), Paragraph(caption, styles["caption"]), Spacer(1, 0.3 * cm)])]
    except Exception:  # ReportLab may reject corrupt or unsupported image data.
        return [Paragraph(f"{caption}: <i>Image unavailable</i>", styles["body"]), Spacer(1, 0.2 * cm)]


def _draw_page_decorations(canvas: Any, document: Any) -> None:
    """Draw the parchment page, disclaimer, and page number around Platypus content."""
    canvas.saveState()
    canvas.setFillColor(PARCHMENT)
    canvas.rect(0, 0, A4[0], A4[1], fill=1, stroke=0)
    canvas.setStrokeColor(BURNT_ORANGE)
    canvas.setLineWidth(0.5)
    canvas.line(1.5 * cm, 1.25 * cm, A4[0] - 1.5 * cm, 1.25 * cm)
    canvas.setFillColor(INK)
    canvas.setFont("Helvetica", 7.5)
    canvas.drawString(1.5 * cm, 0.8 * cm, "Sanchayak report — automated screening aid; final compliance decisions require human review.")
    canvas.drawRightString(A4[0] - 1.5 * cm, 0.8 * cm, f"Page {document.page}")
    canvas.restoreState()


def generate_pdf_report(scan_result: dict) -> bytes:
    """Generate a formatted Legal Metrology PDF report from a scan-result dictionary."""
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.7 * cm, title="Sanchayak Compliance Report")
    styles = _report_styles()
    product = _display(scan_result.get("product_name"), "Unnamed product")
    scan_id = _display(scan_result.get("scan_id"), "—")
    scan_date = _display(scan_result.get("scan_date"), "—")
    verdict = _display(scan_result.get("overall_verdict"), "needs_review").replace("_", " ").upper()
    # Both lookups intentionally fall back to AMBER, a valid ReportLab Color.
    verdict_colour = VERDICT_COLOURS.get(verdict, AMBER)

    story: list[Any] = [
        Paragraph("Sanchayak — Legal Metrology Compliance Report", styles["title"]),
        Paragraph(f"<b>Product:</b> {product}<br/><b>Scan ID:</b> {scan_id}<br/><b>Scan date:</b> {scan_date}", styles["meta"]),
        Spacer(1, 0.35 * cm),
        Table(
            [[Paragraph(verdict, styles["verdict"])]],
            colWidths=[17.5 * cm],
            style=TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), verdict_colour),
                    # ReportLab line commands require width before the Color.
                    ("BOX", (0, 0), (-1, -1), 1, verdict_colour),
                    ("TOPPADDING", (0, 0), (-1, -1), 9),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
                ]
            ),
        ),
        Paragraph("Label evidence", styles["heading"]),
    ]
    story.extend(_image_story(scan_result.get("image_url"), "Original label photo", styles))
    story.extend(_image_story(scan_result.get("annotated_image_url"), "Annotated label photo", styles))
    story.append(Paragraph("Declaration assessment", styles["heading"]))
    rows: list[list[Any]] = [[Paragraph(text, styles["table_header"]) for text in ("Mandatory declaration", "Extracted value", "Status", "Reason")]]
    for item in _normalised_fields(scan_result):
        status = _display(item.get("status"), "flagged").lower()
        status_colour = _status_colour(status)
        rows.append([
            Paragraph(FIELD_LABELS.get(item.get("field_name"), _display(item.get("field_name"), "Unknown field")), styles["body"]),
            Paragraph(_display(item.get("extracted_value"), "Not detected"), styles["body"]),
            Paragraph(
                f'<font color="{status_colour.hexval()}"><b>{status.upper()}</b></font>',
                styles["body"],
            ),
            Paragraph(_display(item.get("reason"), "—"), styles["body"]),
        ])
    assessment_style = TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), INK),
            ("GRID", (0, 0), (-1, -1), 0.35, TABLE_GRID),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BACKGROUND", (0, 1), (-1, -1), TABLE_ROW),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]
    )
    story.append(
        Table(
            rows,
            colWidths=[4.35 * cm, 4.1 * cm, 2.0 * cm, 7.05 * cm],
            repeatRows=1,
            style=assessment_style,
        )
    )
    notes = scan_result.get("inspector_notes")
    if notes:
        story.extend([Paragraph("Inspector notes", styles["heading"]), Paragraph(_display(notes, "—"), styles["body"])])
    doc.build(story, onFirstPage=_draw_page_decorations, onLaterPages=_draw_page_decorations)
    return output.getvalue()


def generate_docx_report(scan_result: dict) -> bytes:
    """Generate an editable DOCX version of the same Sanchayak compliance report."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.65)
    title = document.add_heading("Sanchayak — Legal Metrology Compliance Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    product = _display(scan_result.get("product_name"), "Unnamed product")
    document.add_paragraph(f"Product: {product}\nScan ID: {_display(scan_result.get('scan_id'), '—')}\nScan date: {_display(scan_result.get('scan_date'), '—')}", style="Subtitle")
    verdict = _display(scan_result.get("overall_verdict"), "needs_review").replace("_", " ").upper()
    verdict_paragraph = document.add_paragraph()
    verdict_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    verdict_run = verdict_paragraph.add_run(verdict)
    verdict_run.bold, verdict_run.font.size = True, Pt(16)
    colour = {"COMPLIANT": "5C7A5C", "NON COMPLIANT": "B4483A", "NEEDS REVIEW": "C9922A"}.get(verdict, "C9922A")
    verdict_run.font.color.rgb = RGBColor.from_string(colour)
    for source, caption in ((scan_result.get("image_url"), "Original label photo"), (scan_result.get("annotated_image_url"), "Annotated label photo")):
        document.add_heading(caption, level=2)
        image_data = _read_image(source)
        if image_data:
            try:
                document.add_picture(image_data, width=Inches(6.1))
            except Exception:
                document.add_paragraph("Image unavailable")
        else:
            document.add_paragraph("Image unavailable")
    document.add_heading("Declaration assessment", level=2)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Mandatory declaration", "Extracted value", "Status", "Reason")):
        cell.text = text
    for item in _normalised_fields(scan_result):
        cells = table.add_row().cells
        cells[0].text = FIELD_LABELS.get(item.get("field_name"), _display(item.get("field_name"), "Unknown field"))
        cells[1].text = _display(item.get("extracted_value"), "Not detected")
        cells[2].text = _display(item.get("status"), "flagged").upper()
        cells[3].text = _display(item.get("reason"), "—")
    if scan_result.get("inspector_notes"):
        document.add_heading("Inspector notes", level=2)
        document.add_paragraph(_display(scan_result.get("inspector_notes"), "—"))
    footer = section.footer.paragraphs[0]
    footer.text = "Sanchayak report — automated screening aid; final compliance decisions require human review."
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    result = BytesIO()
    document.save(result)
    return result.getvalue()


def dummy_scan_result(scan_id: str = "demo-scan-001") -> dict[str, Any]:
    """Return representative development data until the database adapter is connected."""
    return {"product_name": "Sanchayak Demo Pack", "scan_id": scan_id, "scan_date": datetime.now(timezone.utc).isoformat(), "image_url": "trap_label.jpg", "annotated_image_url": "trap_label.jpg", "overall_verdict": "needs_review", "fields": [{"field_name": name, "extracted_value": None if name == "country_of_origin" else "Demo extracted value", "status": "flagged" if name == "country_of_origin" else "pass", "reason": "Not detected on label; confirm applicability." if name == "country_of_origin" else None, "font_size_px": None} for name in FIELD_ORDER], "inspector_notes": "Generated from stub data. Replace get_scan_result with the database adapter."}


if __name__ == "__main__":
    Path("compliance_report_demo.pdf").write_bytes(generate_pdf_report(dummy_scan_result()))
    print("Wrote compliance_report_demo.pdf")
